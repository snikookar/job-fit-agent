"""Word export with python-docx: tailored resume highlights + cover letter.

If `template` points to a .docx file, its styles, page setup, header and footer are kept
and the content is appended. That way you can use your own resume design.
With `evidence_comments=True`, each sentence gets a Word comment listing its sources:
handy for a final check, but remove them before sending the file.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

from jobfit.schemas import JOB_EVIDENCE_ID, Chunk, Draft, JobRequirements

GREETINGS = {
    "de": ("Sehr geehrte Damen und Herren,", "Mit freundlichen Grüßen"),
    "en": ("Dear Hiring Team,", "Kind regards,"),
}
HEADINGS = {
    "de": ("Relevante Erfahrung", "Anschreiben"),
    "en": ("Relevant experience", "Cover letter"),
}


def _paragraphs(sentences: list, per_paragraph: int = 3) -> list[list]:
    return [sentences[i : i + per_paragraph] for i in range(0, len(sentences), per_paragraph)]


def build_docx(
    draft: Draft,
    job: JobRequirements,
    profile: dict[str, str],
    path: Path,
    language: str = "en",
    template: Path | None = None,
    evidence: dict[str, Chunk] | None = None,
    evidence_comments: bool = False,
) -> Path:
    doc = Document(str(template)) if template and template.exists() else Document()
    if not template:
        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)

    lang = language if language in GREETINGS else "en"
    greeting, closing = GREETINGS[lang]
    h_experience, h_letter = HEADINGS[lang]
    name = profile.get("name", "")

    def comment(runs, ids: list[str]) -> None:
        if not (evidence_comments and evidence):
            return
        lines = []
        for i in ids:
            if i == JOB_EVIDENCE_ID:
                lines.append("[job posting]")
            elif i in evidence:
                lines.append(f"[{i}] {evidence[i].text}")
        if lines:
            doc.add_comment(runs, text="\n".join(lines), author="job-fit-agent", initials="JFA")

    # --- header
    if name:
        doc.add_heading(name, level=0)
    contact = " · ".join(v for k in ("headline", "email", "phone", "location", "links") if (v := profile.get(k)))
    if contact:
        p = doc.add_paragraph(contact)
        for run in p.runs:
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # --- tailored bullets
    target = f"{job.title}" + (f" — {job.company}" if job.company else "")
    doc.add_heading(f"{h_experience}: {target}", level=1)
    for s in draft.bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(s.text)
        comment([run], s.evidence_ids)

    # --- cover letter on a new page
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading(h_letter, level=1)
    doc.add_paragraph(date.today().strftime("%d.%m.%Y" if lang == "de" else "%d %B %Y"))
    doc.add_paragraph(greeting)
    for group in _paragraphs(draft.cover_letter):
        p = doc.add_paragraph()
        for j, s in enumerate(group):
            run = p.add_run(("" if j == 0 else " ") + s.text)
            comment([run], s.evidence_ids)
    doc.add_paragraph(closing)
    if name:
        doc.add_paragraph(name)

    doc.core_properties.title = f"Application: {target}"
    doc.core_properties.author = name or "job-fit-agent"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
